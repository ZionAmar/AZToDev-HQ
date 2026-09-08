# -*- coding: utf-8 -*-
"""Generate RTL Hebrew Word doc — founder investigation."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_rtl(p)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return p


def add_p(doc, text, bold=False):
    p = doc.add_paragraph()
    set_rtl(p)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    set_rtl(p)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(12)
        r2 = p.add_run(text)
        r2.font.name = "Arial"
        r2.font.size = Pt(12)
    else:
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
    return p


def main():
    out = Path(__file__).resolve().parents[2] / "ops" / "outbox-founder" / "2026-09-07_תחקיר-שיחה-EMET.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    add_heading(doc, "תחקיר מעמיק — שיחת EMET / נעה (CEO)", 0)
    add_p(doc, "מיועד ל: ציון (Founder) | aztodev@gmail.com", bold=True)
    add_p(doc, "תאריך: 7 בספטמבר 2026")
    add_p(doc, "מקור: שיחת Telegram ↔ נעה (Cursor) + מסמכי חברה + KidNest audit")
    doc.add_paragraph()

    # 1
    add_heading(doc, "1. תקציר מנהלים (Executive Summary)", 1)
    add_p(doc, "בשיחה הזאת ביקשת מספר דברים שונים — מ-bootstrap של חברת EMET, דרך דפי נחיתה, ניהול משימות ב-Linear, ועד audit מעמיק של KidNest (NesTube). חלק מהדרך עבד; חלק נכשל בגלל «theater» — הבטחות בלי ביצוע.", bold=True)
    add_bullet(doc, " — EMET עלה, Telegram+נעה עובדים, Linear מחובר ו-25 משימות KidNest עלו.", bold_prefix="מה עובד:")
    add_bullet(doc, " — KidNest audit: 5 מתוך 25 משימות הושלמו; production חי.", bold_prefix="התקדמות אמיתית:")
    add_bullet(doc, " — Linear הובטח לפני publish; delegate נתקע; רק רועי קיבל משימה אחת; secrets נשלחו בטלגרם.", bold_prefix="מה נשבר:")
    add_bullet(doc, " — חברה «חיה» כמו תום אבן, ניהול מהטלפון, KidNest מסודר, בלי CMD שמקריס.", bold_prefix="מה אתה רוצה:")
    doc.add_paragraph()

    # 2
    add_heading(doc, "2. כרונולוגיה — מה קרה בשיחה (לפי נושאים)", 1)

    add_heading(doc, "2.1 Bootstrap ו«מה עכשיו?»", 2)
    add_p(doc, "פתחת שיחה חדשה. EMET היה במצב bootstrap — WIP ריק, מחכה ל-IDEA ראשון. לא היה ברור מה הפעולה הבאה.")
    add_bullet(doc, "הסברנו: אתה זורק IDEA → Intake → Shape → רק אז build.")

    add_heading(doc, "2.2 דפי נחיתה (2 בקשות)", 2)
    add_p(doc, "ביקשת דף נחיתה mobile-first + עדכונים + לינק production; וגרסת עורך דין.", bold=True)
    add_bullet(doc, " — 4 IDEAs נקלטו; intake meetings רצו; build נעצר על CHOOSE (A/B/C).")
    add_bullet(doc, " — Linear ריק; אין לינקים; לא הסברנו פשוט.")
    add_bullet(doc, " — «נוותר על 2 דפי הנחיתה» → killed. לא נבנה כלום.")
    add_p(doc, "מסקנה: הבקשה טופלה ב-process, לא ב-execution. זו «activity, not progress».", bold=True)

    add_heading(doc, "2.3 KidNest — הבקשה האמיתית ש«נשכחה»", 2)
    add_p(doc, "בהקלטה קולית ביקשת: לעבור על KidNest (YouTube מסונן לילדים), desktop+production תקוע, רשימת מטלות, תובנות, דוח.", bold=True)
    add_p(doc, "מה קרה אחר כך: נענה «שולחים לינקים» — **בלי לשלוח**. Linear ריק. Theater.")
    add_p(doc, "רק כששאלת שוב — זיהינו KidNest כ-WIP האמיתי.")

    add_heading(doc, "2.4 Linear — OAuth, MCP, API Key", 2)
    add_bullet(doc, "חיברת Linear ל-Cursor (OAuth ירוק).")
    add_bullet(doc, "MCP עבר מ-Settings ל-Customize → MCPs (Cursor חדש).")
    add_bullet(doc, "OAuth ≠ publish. אף issue לא נוצר עד API key ב-.env.")
    add_bullet(doc, "7/9/2026: 25 משימות KidNest עלו ל-team EMET ב-Linear.")
    add_p(doc, "לקח: OAuth מספיק ל-Cursor; publish אוטומטי דורש API key או MCP שנטען ב-session.", bold=True)

    add_heading(doc, "2.5 KidNest Audit — ביצוע", 2)
    add_p(doc, "25 משימות (KN-001…KN-025) ב-5 שלבים. נכון לסוף השיחה:", bold=True)
    add_bullet(doc, "KN-001 — מיפוי 94 API routes")
    add_bullet(doc, "KN-004 — 27 ערוצים קיבלו קטגוריה (0 נמחקו — לפי הוראתך)")
    add_bullet(doc, "KN-008 — player: קוד נראה תקין; לא נבדק runtime")
    add_bullet(doc, "KN-014 — prod env: רוב SET; YouTube/Google secret לבדוק ב-cPanel")
    add_bullet(doc, "KN-015 — production חי: health/login/APK PASS")
    add_p(doc, "20 משימות נשארו. אין דוח סופי (KN-025) עדיין.")

    add_heading(doc, "2.6 CMD, כיבוי מחשב, תפעול", 2)
    add_bullet(doc, "CMD יתומים — ניקוי כל 12 שניות (EMET runtime) + אחרי delegates.")
    add_bullet(doc, "ביקשת כיבוי מלא — בוצע (shutdown /s /t 10).")
    add_bullet(doc, "EMET-OFF / EMET-ON — לריסטart runtime.")

    doc.add_paragraph()

    # 3 Tom Even
    add_heading(doc, "3. תום אבן (Tom Even / agents&me) — מה זה ומה הקשר", 1)
    add_p(doc, "תום אבן = יוצר/מנחה סדנאות AI Agents בישראל (agents&me, ABC-TOM). מערכת של ~12 סוכנים על המחשב, knowledge base משותף, 6 תיקיות, עבודה «בזמן שישן».", bold=True)
    add_p(doc, "קישור Gemini ששלחת לא נפתח (דורש login). EMET מגדירה «Tom Even Parity+» במסמך _company/TOM_EVEN_PARITY.md.")

    add_heading(doc, "3.1 מה תום אבן נותן", 2)
    add_bullet(doc, "סוכנים עם זהות (CEO, copywriter, analyst…)")
    add_bullet(doc, "שיחה בין סוכנים (קבצים / Obsidian)")
    add_bullet(doc, "אוטומציות יומיות, דיילי")
    add_bullet(doc, "תחושת «חברה חיה»")
    add_bullet(doc, "בסדנה: Claude Code (Cursor אפשרי)")

    add_heading(doc, "3.2 מה EMET מתוכננת לתת (Parity+)", 2)
    add_bullet(doc, " — Telegram ↔ נעה בלבד (לא 30 צ'אטים)", bold_prefix="יותר טוב:")
    add_bullet(doc, " — Linear בטלפון", bold_prefix="יותר טוב:")
    add_bullet(doc, " — Stage Machine, gates, QA", bold_prefix="יותר טוב:")
    add_bullet(doc, " — Cursor מקומי, products/ אמיתי", bold_prefix="יותר טוב:")
    add_bullet(doc, " — ON/OFF + checkpoint", bold_prefix="יותר טוב:")
    add_bullet(doc, " — 33 סוכנים + bus + meetings", bold_prefix="שווה:")
    add_bullet(doc, " — emet_delegate = session נפרד לכל סוכן", bold_prefix="שווה:")

    add_heading(doc, "3.3 פערים — למה זה לא «כמו תום אבן» עדיין", 2)
    add_bullet(doc, "Orchestrator לא מריץ daily/intake אוטומטית עם LLM באופן עקבי", bold_prefix="❌")
    add_bullet(doc, "Dispatch — משימות ב-Linear ≠ סוכנים שעובדים", bold_prefix="❌")
    add_bullet(doc, "Theater — תשובות לפני tool", bold_prefix="❌")
    add_bullet(doc, "Linear MCP לא נטען ב-sessions (רק API key פתר)", bold_prefix="❌")
    add_bullet(doc, "סיכום EOD אוטומטי — חלקי", bold_prefix="⚠️")
    add_bullet(doc, "KidNest audit התחיל רק אחרי לחץ חוזר", bold_prefix="⚠️")

    doc.add_paragraph()

    # 4 What founder wants
    add_heading(doc, "4. מה אתה רוצה (מסכמים מהשיחה)", 1)
    add_bullet(doc, "חברת AI agents שעובדת באמת — לא theater")
    add_bullet(doc, "לראות משימות מהטלפון (Linear) — לא dashboard מקומי")
    add_bullet(doc, "KidNest מובן + מתוקן + דוח")
    add_bullet(doc, "סוכנים עצמאיים ב-Cursor, מעבירים תוצרים, מתקשרים ב-bus")
    add_bullet(doc, "עדכונים קצרים מדי פעם — לא הצפה")
    add_bullet(doc, "בלי CMD שמקריס את המחשב")
    add_bullet(doc, "שיחה טבעית — תמונה, קול, קבצים (חלקית מיושם)")

    doc.add_paragraph()

    # 5 Recommendations
    add_heading(doc, "5. איך הכי כדאי לך לעשות את מה שאתה רוצה — המלצות", 1)

    add_heading(doc, "5.1 חוזה תפעולי (מה לדרוש מנעה)", 2)
    add_p(doc, "כלל ברזל: Think-aloud ≠ execute. אין «שלחתי» / «עובדים» בלי artifact.", bold=True)
    add_bullet(doc, "משימה חדשה → Linear issue + delegate + קובץ תוצר")
    add_bullet(doc, "עדכון אליך: סוף יום, P0 done, שער waiting_founder")
    add_bullet(doc, "WIP max 2–3 משימות במקביל")
    add_bullet(doc, "Secrets רק ב-.env — never Telegram")

    add_heading(doc, "5.2 Linear + טלפון", 2)
    add_p(doc, "שמור LINEAR_API_KEY ב-.env (rotate אחרי ששלחת בטלגרם!).", bold=True)
    add_bullet(doc, "Linear app — login אותו חשבון")
    add_bullet(doc, "Team: EMET — 25 issues KidNest")
    add_bullet(doc, "Filter: label / project KidNest")

    add_heading(doc, "5.3 KidNest — המשך מומלץ", 2)
    add_bullet(doc, "KN-010 — smoke test (player באמת עובד?)")
    add_bullet(doc, "KN-002 — monorepo setup doc")
    add_bullet(doc, "KN-022 — security audit")
    add_bullet(doc, "KN-025 — דוח תובנות בעברית (אימייל + Telegram)")
    add_p(doc, "Production לא «תקוע» — nestube.aztodev.com חי. יש פערים ב-env (YouTube/Google).", bold=True)

    add_heading(doc, "5.4 «כמו תום אבן» — 3 צעדים מעשיים", 2)
    add_p(doc, "1. הפעל EMET-OFF → EMET-ON אחרי שינוי .env", bold=True)
    add_p(doc, "2. דרוש daily אוטומטי + EOD summary — אם לא מגיע, זה bug", bold=True)
    add_p(doc, "3. לכל בקשה: «מה ב-Linear?» + «הראה קובץ» — לא «אנחנו על זה»", bold=True)

    add_heading(doc, "5.5 CMD — הגנה על המחשב", 2)
    add_bullet(doc, "EMET runtime מנקה cmd/conhost כל 12 שניות")
    add_bullet(doc, "אם הצפה: EMET-OFF, סגור cmd ריקים, EMET-ON")
    add_bullet(doc, "EMET_CURSOR_SETTING_SOURCES=project — פחות hooks")
    add_bullet(doc, "אל תריץ 5 delegates במקביל ללא צורך")

    add_heading(doc, "5.6 אבטחה — דחוף", 2)
    add_p(doc, "שלחת LINEAR_API_KEY בטלגרם. מומלץ:", bold=True)
    add_bullet(doc, "Linear → revoke key → create new → .env only")
    add_bullet(doc, "אותו דבר אם נחשף CURSOR/OPENAI key")

    doc.add_paragraph()

    # 6 Status table
    add_heading(doc, "6. סטטוס נוכחי (7/9/2026)", 1)
    add_p(doc, "KidNest Audit: 5/25 done | Linear: 25 issues | Landing: killed | Runtime: ON", bold=True)

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    rows = [
        ("נושא", "סטטוס"),
        ("Telegram ↔ נעה", "✅ פעיל"),
        ("Linear (טלפון)", "✅ 25 issues (אחרי API key)"),
        ("KidNest audit", "⚠️ 5/25 — production חי"),
        ("דפי נחיתה", "❌ killed"),
        ("Parity+ (תום אבן)", "⚠️ חלקי — theater תוקן חלקית"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
        for c in table.rows[i].cells:
            for p in c.paragraphs:
                set_rtl(p)

    doc.add_paragraph()

    # 7 Next
    add_heading(doc, "7. מה לעשות מחר בבוקר (Action List)", 1)
    add_p(doc, "1. Rotate Linear API key", bold=True)
    add_p(doc, "2. פתח Linear → וודא 25 משימות EMET", bold=True)
    add_p(doc, "3. כתוב לנעה: «המשיכי KidNest — smoke + דוח»", bold=True)
    add_p(doc, "4. EMET-OFF → EMET-ON אם שינית .env", bold=True)
    add_p(doc, "5. אל תפתח landing מחדש — KidNest הוא WIP", bold=True)

    add_heading(doc, "8. נספח — KidNest בקצרה", 1)
    add_p(doc, "NesTube/KidNest = YouTube בטוח לילדים. Admin להורים, Mobile/TV לילדים. Production: nestube.aztodev.com. Monorepo: api + admin + mobile + tv-home + tv.")
    add_p(doc, "ממצאים עיקריים: production חי; 27 ערוצים תויגו; player צריך smoke; Google OAuth לא מחובר ב-UI; 8 orphan API routes.")

    add_p(doc, "— נעה (CEO), EMET | Build truth. Ship value. Scale leverage.", bold=True)

    doc.save(str(out))
    print(str(out))


if __name__ == "__main__":
    main()
