# -*- coding: utf-8 -*-
"""Tom Even vs EMET deep comparison — Hebrew RTL Word doc."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_rtl(p)
    for r in p.runs:
        r.font.name = "Arial"
    return p


def p(doc, text, bold=False):
    para = doc.add_paragraph()
    set_rtl(para)
    r = para.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.bold = bold
    para.paragraph_format.space_after = Pt(6)
    return para


def bullet(doc, prefix, text):
    para = doc.add_paragraph(style="List Bullet")
    set_rtl(para)
    if prefix:
        r1 = para.add_run(prefix)
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(12)
    r2 = para.add_run(text)
    r2.font.name = "Arial"
    r2.font.size = Pt(12)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, hd in enumerate(headers):
        t.rows[0].cells[i].text = hd
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = val
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                set_rtl(para)
                for r in para.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)


def main():
    out = Path(__file__).resolve().parents[2] / "ops" / "outbox-founder" / "2026-09-07_Tom-Even-vs-EMET-deep-he.docx"
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)

    h(doc, "השוואה מעמיקה: תום אבן (agents&me) מול EMET", 0)
    p(doc, "מיועד ל: ציון | תוספת לתחקיר השיחה", bold=True)
    p(doc, "7 בספטמבר 2026 — נעה (CEO)")
    doc.add_paragraph()

    h(doc, "1. במה תום אבן משתמש בפועל?", 1)
    p(doc, "תום אבן = מנחה סדנאות AI Agents (agents&me). הוא לא מוכר «מוצר SaaS» — הוא מלמד שיטה (ABC-TOM) לבנות צוות סוכנים על המחשב שלך.", bold=True)

    h(doc, "1.1 Stack טכני (לפי אתר + LinkedIn + newsletter)", 2)
    bullet(doc, "Claude Code (CLI): ", "הכלי הראשי היומי. טerminal שמדבר עם Claude, קורא/כותב קבצים, מריץ פקודות. בסדנה: «Claude Code — הסטאק שאני משתמש בו כל יום».")
    bullet(doc, "מנוי Anthropic: ", "Claude Pro ~$20/חודש, או Claude Max ~$100/חודש (תום על Max — יותר quota ל-Opus/Sonnet).")
    bullet(doc, "מודלים: ", "Claude Opus / Sonnet / Haiku — דרך Claude Code. Opus = חשיבה עמוקה; Sonnet = איזון; Haiku = מהיר/זול. אין «Composer» — הכל Anthropic.")
    bullet(doc, "ABC-TOM: ", "מערכת הפעלה שהוא המציא — 6 תיקיות + קבצים = «מוח שני» + כללי עבודה לסוכנים.")
    bullet(doc, "קבצים לוקאליים: ", "Markdown / notes — זיכרון משותף. לא Obsidian חובה, אבל רוח דומה.")
    bullet(doc, "~12 סוכנים: ", "CEO, copywriter, designer, analyst, CTO, marketing… כל אחד = prompt + תיקייה + תפקיד.")
    bullet(doc, "Cursor: ", "«אפשר גם Cursor אם ממש רוצים» — לא הליבה. Claude Code חזק יותר ל-terminal/file workflow לדבריו.")
    bullet(doc, "Neo: ", "סוכן «אוטונומי» שעושה דברים מעניינים — לא מלמדים איך בנוי.")
    bullet(doc, "אין (בסדנה): ", "Telegram bot, Node runtime, Linear, stage gates, deploy production.")

    h(doc, "1.2 מה הוא מייצר?", 2)
    bullet(doc, "", "תוכן: newsletter, LinkedIn, copy, מצגות, דוחות")
    bullet(doc, "", "אופרציה: מחקר מתחרים, סיכומי ראיונות, מעקב משימות")
    bullet(doc, "", "אתרים/דפים (דוגמה: האתר שלו ב-56 דקות)")
    bullet(doc, "", "פחות: SaaS production, QA gates, security review, monorepo engineering")

    doc.add_paragraph()

    h(doc, "2. במה EMET (השיטה שלך) משתמשת?", 1)
    bullet(doc, "Cursor IDE + Cursor SDK: ", "CURSOR_API_KEY — מריץ סוכנים מחוץ לחלון (Telegram, orchestrator). מודל: Composer 2.5 (ברירת מחדל).")
    bullet(doc, "OpenAI API: ", "gpt-4.1-mini — שיחת Telegram מהירה עם נעה (לא Cursor לכל הודעה).")
    bullet(doc, "Node.js runtime (EMET-ON): ", "emet.mjs — Telegram, ticks, daily, intake meetings, CMD cleanup.")
    bullet(doc, "33 סוכנים: ", "תיקייה לכל role — SYSTEM_PROMPT, PERSONALITY, PERMISSIONS.")
    bullet(doc, "emet_delegate: ", "פותח chat Cursor נפרד לכל מומחה (PM, CTO, Tech Lead…).")
    bullet(doc, "ops/bus + ops/state.json: ", "תקשורת בין סוכנים — כמו «חברה».")
    bullet(doc, "Linear + API key: ", "משימות בטלפון.")
    bullet(doc, "Telegram ↔ נעה בלבד: ", "Front desk — לא 30 צ'אטים.")
    bullet(doc, "Stage Machine: ", "Intake → Discover → Shape → Build → Launch + gates.")
    bullet(doc, "products/: ", "קוד אמיתי — KidNest, work_clock, וכו'.")
    bullet(doc, "MCP: ", "Linear, custom tools (emet_delegate, status).")

    doc.add_paragraph()

    h(doc, "3. טבלת השוואה — יכולות", 1)
    add_table(doc, ["נושא", "תום אבן", "EMET", "מי מנצח"], [
        ["IDE / Agent", "Claude Code (CLI)", "Cursor + SDK", "שונה — לא «מי טוב»"],
        ["מודל AI", "Claude Opus/Sonnet (Max)", "Composer 2.5 + GPT-4.1-mini", "תלוי משימה"],
        ["זיכרון", "6 תיקיות ABC-TOM", "ops/ + agents/ + SHARED.md", "EMET — יותר structured"],
        ["מספר סוכנים", "~12", "33", "EMET"],
        ["תקשורת ביניהם", "קבצים משותפים", "bus + meetings + delegate", "EMET — יותר «חברה»"],
        ["ערוץ אליך", "אתה ב-Claude Code", "Telegram ↔ נעה", "EMET — טלפון"],
        ["משימות מרחוק", "קבצים / Obsidian-like", "Linear app", "EMET"],
        ["קוד / SaaS", "אפשרי, לא מוקד", "מוקד ראשי", "EMET"],
        ["Production gates", "לא", "כן (APPROVE)", "EMET"],
        ["אוטומציה ברקע", "ידני / scripts", "EMET-ON runtime", "EMET"],
        ["עלות חודשית", "$20–100 Claude", "Cursor Pro + API keys", "דומה"],
        ["פשטות setup", "ZIP + 2.5h workshop", "repo + .env + EMET-ON", "תום — פשוט יותר"],
        ["תוכן / marketing", "מוקד חזק", "CMO agents — לא הופעל", "תום — היום"],
        ["Engineering audit", "לא מוגדר", "KidNest KN-*", "EMET"],
    ])

    doc.add_paragraph()

    h(doc, "4. מה תום יכול לעשות — ואתה EMET לא (עדיין)", 1)
    bullet(doc, "Claude Code terminal flow: ", "שיחה + קבצים + shell בזרימה אחת — מאוד חלק ל-content creators. Cursor CLI/SDK פחות «טבעי» ל-non-devs.")
    bullet(doc, "ABC-TOM מוכן: ", "ZIP + workshop = עובד ביום 1. EMET דורש bootstrap + .env + למידה.")
    bullet(doc, "מהירות תוכן: ", "newsletter, posts, copy — Opus + Claude Code מאומץ לזה. EMET מפנה ל-gates לפני publish.")
    bullet(doc, "פשטות נפשית: ", "אתה מדבר עם Claude ישירות. EMET = שכבת נעה + runtime (יותר עוצמה, יותר complexity).")
    bullet(doc, "קהילת בוגרים: ", "600+ alumni — תמיכה peer. EMET = solo founder + agents.")
    bullet(doc, "עלות צפויה: ", "Claude Max $100 — הכל inclusive. EMET = Cursor + OpenAI + Linear (מפוזר).")

    h(doc, "5. מה EMET יכולה — ותום לא מדגיש", 1)
    bullet(doc, "SaaS אמיתי: ", "KidNest audit — 94 routes, prod health, 27 channels tagged. קוד ב-repo.")
    bullet(doc, "Telegram front desk: ", "אתה מהטלפון — לא צריך לפתוח IDE.")
    bullet(doc, "Linear בטלפון: ", "25 issues — רואה WIP מרחוק.")
    bullet(doc, "Stage gates: ", "לא deploy / לא spend בלי APPROVE.")
    bullet(doc, "33 roles + permissions: ", "CFO לא נוגע ב-production. Least privilege.")
    bullet(doc, "EMET-OFF/ON + checkpoint: ", "כיבוי מסודר + resume.")
    bullet(doc, "WIP limits: ", "קים — לא 25 משימות במקביל.")
    bullet(doc, "CMD cleanup: ", "ניקוי cmd יתומים — תום לא מדבר על זה.")
    bullet(doc, "Voice/media Telegram: ", "Whisper, תמונות, reply threading.")

    doc.add_paragraph()

    h(doc, "6. Claude Opus / Sonnet / Haiku — מה זה אומר לך?", 1)
    p(doc, "Claude Code = ממשק. Opus/Sonnet/Haiku = «מוח» בתוך Anthropic.", bold=True)
    bullet(doc, "Opus: ", "הכי חכם — ארכיטקטורה, ניתוח עמוק, כתיבה ארוכה. יקר, איטי יותר.")
    bullet(doc, "Sonnet: ", "יום-יום — 80% מהמשימות.")
    bullet(doc, "Haiku: ", "מהיר/זול — סיכומים, tagging.")
    p(doc, "ב-EMET המקבילה:", bold=True)
    bullet(doc, "Composer 2.5 (Cursor): ", "build + refactor + multi-file — כמו Sonnet/Opus hybrid.")
    bullet(doc, "gpt-4.1-mini: ", "Telegram chat מהיר — כמו Haiku.")
    bullet(doc, "gpt-4.1 smart: ", "ישיבות / החלטות — כמו Sonnet.")
    p(doc, "אין לך Claude Opus ישירות ב-EMET — אלא אם תוסיף ANTHROPIC_API_KEY או תריץ Claude Code במקביל.", bold=True)

    doc.add_paragraph()

    h(doc, "7. יתרונות וחסרונות — לפי «מה אתה רוצה»", 1)

    h(doc, "7.1 אם המטרה: «חברה שבונה SaaS (KidNest…)»", 2)
    p(doc, "EMET — יתרון. תום — חסרון.", bold=True)
    bullet(doc, "EMET +: ", "gates, QA, deploy control, 33 roles, real repo.")
    bullet(doc, "תום +: ", "מהיר להתחיל content, פחות overhead.")
    bullet(doc, "סיכון EMET: ", "theater — planning בלי execution (ראינו בשיחה).")
    bullet(doc, "סיכון תום: ", "אין production discipline — קל «לשכוח» deploy/security.")

    h(doc, "7.2 אם המטרה: «לנהל מהטלפון בלי IDE»", 2)
    p(doc, "EMET — יתרון (Telegram + Linear).", bold=True)
    bullet(doc, "תום: ", "אתה ב-Claude Code על המחשב — לא phone-first.")

    h(doc, "7.3 אם המטרה: «תוכן, marketing, newsletter»", 2)
    p(doc, "תום — יתרון היום.", bold=True)
    bullet(doc, "EMET: ", "יש CMO agents — לא הופעלו. צריך delegate + brief.")

    h(doc, "7.4 אם המטרה: «פשוט, עובד מחר»", 2)
    p(doc, "תום — יתרון (workshop + ZIP).", bold=True)
    bullet(doc, "EMET: ", "מורכב — אבל scalable ל-SaaS.")

    doc.add_paragraph()

    h(doc, "8. המלצה: איך לקבל את הטוב משניהם", 1)
    p(doc, "אל תזרוק EMET. אל תעתיק תום 1:1. Hybrid:", bold=True)
    bullet(doc, "1. ", "שמור EMET ל-SaaS + KidNest + gates + Telegram + Linear.")
    bullet(doc, "2. ", "הוסף Claude Code ל-content/marketing (newsletter, copy) — $20 Pro מספיק להתחלה.")
    bullet(doc, "3. ", "דרוש מנעה: delegate אמיתי — לא theater. WIP max 3.")
    bullet(doc, "4. ", "אופציונלי: ANTHROPIC_API_KEY ב-.env לישיבות orchestrator (Opus-level).")
    bullet(doc, "5. ", "אל תפתח 25 agents במקביל — זה CMD + noise.")
    bullet(doc, "6. ", "Rotate secrets — Linear key נחשף בטלגרם.")

    h(doc, "9. תשובה ישירה לשאלותיך", 1)
    p(doc, "«במה הוא משתמש?» — Claude Code + Claude Max + ABC-TOM + 6 folders + ~12 agents.", bold=True)
    p(doc, "«Claude Opus?» — כן, דרך Max subscription ב-Claude Code. לא Cursor.", bold=True)
    p(doc, "«מה הוא יכול ואני לא?» — content velocity, simpler onboarding, terminal-native AI, alumni community.", bold=True)
    p(doc, "«מה אני יכול והוא לא?» — SaaS engineering, phone management, Linear, production gates, 33-role company.", bold=True)
    p(doc, "«מה הכי כדאי?» — EMET ל-product; שקול Claude Code ל-content; תקן dispatch + theater.", bold=True)

    p(doc, "— נעה, EMET", bold=True)
    doc.save(str(out))
    print(str(out))


if __name__ == "__main__":
    main()
